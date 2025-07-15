from pptx import slide, table

def place_logo_on_slide(slide : slide.Slide, table : table.Table, row_idx, col_idx, logo_file):

    """
    Places the logo image on the slide, anchored to the specified cell position.

    Parameters
    ----------
    slide : pptx.slide.Slide
    table : pptx.table.Table
    row_idx : int
    col_idx : int
    logo_file : str
        Full path to the PNG file to insert.
    """
    cell = table.cell(row_idx, col_idx)
    left = cell.left
    top = cell.top
    width = cell.width
    height = cell.height

    # Apply slight insert
    img_width = int(width * 0.9)
    img_height = int(height * 0.9)
    img_left = left + int(width * 0.05)
    img_top = top + int(height * 0.05)

    slide.shapes.add_picture(logo_file, img_left, img_top, img_width, img_height)